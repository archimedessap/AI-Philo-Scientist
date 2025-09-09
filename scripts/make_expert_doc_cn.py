#!/usr/bin/env python3
import csv
import json
from pathlib import Path
from datetime import datetime
from make_expert_doc import latest_top10_dir, load_summary, extract_texts


def try_make_docx(cards: list, out_path: Path) -> bool:
    try:
        from docx import Document
        from docx.shared import Pt
    except Exception:
        return False

    doc = Document()
    doc.add_heading('前10个生成量子理论 — 专家评审包（中文）', level=0)
    p = doc.add_paragraph()
    p.add_run('生成时间：').bold = True
    p.add_run(datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
    doc.add_paragraph('使用说明：每个理论卡片包含摘要、关键要点与评分项，请逐项阅读并给出评分与评语。').italic = True

    for i, c in enumerate(cards, 1):
        doc.add_page_break()
        doc.add_heading(f"#{i} {c['name']}", level=1)
        if c.get('model_src') or c.get('model_name'):
            doc.add_paragraph(f"模型：{c.get('model_src','')}/{c.get('model_name','')}")

        if c.get('summary'):
            doc.add_heading('摘要', level=2)
            doc.add_paragraph(c['summary'])

        if c.get('core_principles'):
            doc.add_heading('核心原则', level=2)
            for cp in c['core_principles']:
                doc.add_paragraph(cp, style='List Bullet')

        if c.get('formalism'):
            doc.add_heading('形式化（要点）', level=2)
            doc.add_paragraph(c['formalism'])

        if c.get('predictions'):
            doc.add_heading('预测与可验证性', level=2)
            for pr in c['predictions']:
                doc.add_paragraph(pr, style='List Number')

        if c.get('measurement'):
            doc.add_heading('测量处理 / 哲学（简述）', level=2)
            doc.add_paragraph(c['measurement'])

        # Scoring section (CN)
        doc.add_heading('专家评分（0–10）', level=2)
        for field in ['数学严谨性', '可测试性', '本体清晰度', '内在一致性', '综合评分']:
            run = doc.add_paragraph().add_run(f"{field}：__/10    ")
            run.bold = True
        doc.add_paragraph('评语：')
        doc.add_paragraph('\n' * 2)

    doc.save(out_path)
    return True


def make_rtf(cards: list, out_path: Path):
    lines = []
    lines.append('{\\rtf1\\ansi\\deff0')
    lines.append('\\b 前10个生成量子理论 — 专家评审包（中文）\\b0\\par')
    lines.append(f'生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}\\par')
    lines.append('使用说明：每个理论卡片包含摘要、关键要点与评分项，请逐项阅读并给出评分与评语。\\par')
    lines.append('\\par')
    def esc(s: str) -> str:
        return s.replace('\\', r'\\').replace('{', r'\{').replace('}', r'\}')
    for i, c in enumerate(cards, 1):
        lines.append('\\par')
        lines.append('\\pard____________________________________________________________\\par')
        lines.append('\\pard')
        lines.append(f'\\b #{i} {esc(c["name"])}\\b0\\par')
        if c.get('model_src') or c.get('model_name'):
            lines.append(f'模型：{esc(c.get("model_src",""))}/{esc(c.get("model_name",""))}\\par')
        if c.get('summary'):
            lines.append('\\b 摘要\\b0\\par')
            lines.append(esc(c['summary']) + '\\par')
        if c.get('core_principles'):
            lines.append('\\b 核心原则\\b0\\par')
            for cp in c['core_principles']:
                lines.append('• ' + esc(cp) + '\\par')
        if c.get('formalism'):
            lines.append('\\b 形式化（要点）\\b0\\par')
            lines.append(esc(c['formalism']) + '\\par')
        if c.get('predictions'):
            lines.append('\\b 预测与可验证性\\b0\\par')
            for pr in c['predictions']:
                lines.append('- ' + esc(pr) + '\\par')
        if c.get('measurement'):
            lines.append('\\b 测量处理 / 哲学（简述）\\b0\\par')
            lines.append(esc(c['measurement']) + '\\par')
        lines.append('\\b 专家评分（0–10）\\b0\\par')
        for f in ['数学严谨性', '可测试性', '本体清晰度', '内在一致性', '综合评分']:
            lines.append(f'{f}：__/10\\par')
        lines.append('评语：\\par\\par')
    lines.append('}')
    out_path.write_text('\n'.join(lines), encoding='utf-8')


def main():
    base = Path('expert_scoring')
    topdir = latest_top10_dir(base)
    if not topdir:
        print('未找到 expert_scoring/top10_* 目录，请先导出前10个理论。')
        return
    summary = load_summary(topdir)
    if not summary:
        print('未找到 summary.csv：', topdir)
        return

    cards = []
    for row in summary:
        jf = topdir / row['file']
        if not jf.exists():
            jf = Path(topdir) / Path(row['file']).name
        try:
            obj = json.loads(jf.read_text(encoding='utf-8'))
        except Exception:
            continue
        c = extract_texts(obj)
        c['model_src'] = row.get('model_src') or c.get('model_src')
        c['model_name'] = row.get('model_name') or c.get('model_name')
        cards.append(c)

    out_docx = topdir / 'top10_expert_review_CN.docx'
    out_rtf = topdir / 'top10_expert_review_CN.rtf'

    if try_make_docx(cards, out_docx):
        print('已生成：', out_docx)
    else:
        make_rtf(cards, out_rtf)
        print('已生成：', out_rtf)


if __name__ == '__main__':
    main()

